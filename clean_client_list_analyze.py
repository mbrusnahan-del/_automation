"""
clean_client_list_analyze.py
============================

STANDALONE analyzer — dry-run only. Does NOT modify Master-Business Plan.xlsm.

Reads the 'Client List' sheet and flags contact-cell issues to support the
Notion sync cleanup:

  1. Multi-person cells       → proposed split into adjacent Contact columns
  2. Trailing junk            → credentials (PE, AIA, CFM), labels (Cell,
                                Ext. 201, Project Manager), stray punctuation
                                (trailing apostrophes, commas)
  3. Broken emails            → flagged for human review (e.g. jeremy@jeremy@...)
  4. Suspicious long names    → 4+ word names with no separator — possibly
                                two people smushed together (manual review)
  5. Duplicates within client → same name in multiple Contact columns for
                                the same Client row

Outputs: client_list_cleanup_report.xlsx (same folder as this script)

Usage:
  python clean_client_list_analyze.py
  python clean_client_list_analyze.py --excel "C:\\...\\Master-Business Plan.xlsm"
  python clean_client_list_analyze.py --client 42
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from dataclasses import dataclass, field
from typing import Optional

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))


# ── Config ──────────────────────────────────────────────────────────────
CLIENT_LIST_SHEET = "Client List"
DATA_START_ROW    = 4
COL_NUMBER        = "B"
COL_NAME          = "E"
CONTACT_COLS      = ["G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R"]

DEFAULT_EXCEL_PATHS = [
    os.path.join(os.path.dirname(os.path.dirname(_SCRIPT_DIR)),
                 "Business Plan", "Master-Business Plan.xlsm"),
    "/sessions/nice-compassionate-pasteur/mnt/OneDrive - Kingdom Structural LLC/Business Plan/Master-Business Plan.xlsm",
]

OUTPUT_PATH = os.path.join(_SCRIPT_DIR, "client_list_cleanup_report.xlsx")


# ── Regex ───────────────────────────────────────────────────────────────
EMAIL_RE = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")
# Phone pattern — accepts extension like "x2" or "ext 201"
PHONE_RE = re.compile(
    r"(?<!\d)\(?\d{3}\)?[.\-\s]?\d{3}[.\-\s]?\d{4}"
    r"(?:\s*(?:[xX]\.?|ext\.?)\s*\d+)?(?!\d)",
    re.IGNORECASE,
)
BROKEN_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9._%+-]+@")
AMP_SEP_RE = re.compile(r"\s+&\s+")
AND_SEP_RE = re.compile(r"\s+and\s+", re.IGNORECASE)
APOS_TRAIL_RE  = re.compile(r"[\"'`]+$")
PUNCT_TRAIL_RE = re.compile(r"[\s,;:]+$")

# Professional credentials that sometimes get appended to names
CRED_LIST = r"(?:PE|CFM|AIA|SE|RA|CBO|LEED\s*AP|LEED|ASCE|PMP|EIT|NCARB|RLA|PLS|CCM)"
CRED_TAIL_RE = re.compile(
    rf"\s+{CRED_LIST}(?:\s+{CRED_LIST})*\.?\s*$",
    re.IGNORECASE,
)
# Descriptive labels that end up on a name ("… Cell", "… Ext. 201")
LABEL_TAIL_RE = re.compile(
    r"\s+(?:Cell|Mobile|Home|Work|Office|Project\s+Manager|Ext\.?\s*\d+)\s*$",
    re.IGNORECASE,
)


# ── Data classes ────────────────────────────────────────────────────────
@dataclass
class CellIssue:
    row: int
    client_number: str
    client_name: str
    col: str
    original: str
    kind: str                 # "single_fix" | "split" | "manual" | "duplicate"
    fix_type: str
    cleaned: str = ""
    split_parts: list = field(default_factory=list)
    note: str = ""


# ── Helpers ─────────────────────────────────────────────────────────────
def _stringify_number(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def _collapse_ws(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().strip(",;:").strip()


def strip_email_phone(text: str) -> tuple[str, list[str], list[str]]:
    emails = EMAIL_RE.findall(text)
    phones = PHONE_RE.findall(text)
    leftover = text
    for e in emails:
        leftover = leftover.replace(e, " ")
    for p in phones:
        leftover = leftover.replace(p, " ")
    return _collapse_ws(leftover), emails, phones


def apply_tail_strippers(name: str) -> tuple[str, list[str]]:
    """Strip trailing credentials / labels / punct. Returns (cleaned, applied_tags).
    Runs in a loop because a name can have both ('Erica Flannery AIA Project Manager').
    """
    applied: list[str] = []
    changed = True
    while changed:
        changed = False
        m = LABEL_TAIL_RE.search(name)
        if m:
            applied.append(f"label:{m.group(0).strip()}")
            name = LABEL_TAIL_RE.sub("", name)
            changed = True
        m = CRED_TAIL_RE.search(name)
        if m:
            applied.append(f"credentials:{m.group(0).strip()}")
            name = CRED_TAIL_RE.sub("", name)
            changed = True
        if APOS_TRAIL_RE.search(name):
            applied.append("trailing_apostrophe")
            name = APOS_TRAIL_RE.sub("", name)
            changed = True
        if PUNCT_TRAIL_RE.search(name):
            applied.append("trailing_punct")
            name = PUNCT_TRAIL_RE.sub("", name)
            changed = True
    return name.strip(), applied


def split_multi_person(name_part: str,
                       emails: list[str],
                       phones: list[str]) -> Optional[list[dict]]:
    """Detect a multi-person cell and propose a split. Returns None if not multi."""
    # Two emails = two people
    if len(emails) >= 2:
        parts = None
        if AMP_SEP_RE.search(name_part):
            parts = [p.strip() for p in AMP_SEP_RE.split(name_part, maxsplit=1)]
        elif AND_SEP_RE.search(name_part):
            parts = [p.strip() for p in AND_SEP_RE.split(name_part, maxsplit=1)]
        if not parts or len(parts) != 2 or not all(parts):
            # Can't confidently pair names to emails — still flag as split but
            # the apply step will show the raw emails for manual alignment.
            people = []
            for i, e in enumerate(emails):
                p = phones[i] if i < len(phones) else ""
                people.append({"name": name_part if i == 0 else "",
                               "email": e, "phone": p})
            return people
        return [
            {"name": _collapse_ws(parts[0]),
             "email": emails[0] if len(emails) > 0 else "",
             "phone": phones[0] if len(phones) > 0 else ""},
            {"name": _collapse_ws(parts[1]),
             "email": emails[1] if len(emails) > 1 else "",
             "phone": phones[1] if len(phones) > 1 else ""},
        ]

    # Single email (or none) + "&" or " and " separator
    for sep_re in (AMP_SEP_RE, AND_SEP_RE):
        if sep_re.search(name_part):
            parts = [_collapse_ws(p) for p in sep_re.split(name_part, maxsplit=1)]
            if len(parts) == 2 and all(parts):
                return [
                    {"name": parts[0],
                     "email": emails[0] if emails else "",
                     "phone": phones[0] if phones else ""},
                    {"name": parts[1], "email": "", "phone": ""},
                ]
    return None


def analyze_cell(raw, row: int, col: str,
                 client_number: str, client_name: str) -> list[CellIssue]:
    issues: list[CellIssue] = []
    if raw is None:
        return issues
    text = str(raw).replace("\n", " ").strip()
    if not text:
        return issues

    if BROKEN_EMAIL_RE.search(text):
        issues.append(CellIssue(
            row=row, client_number=client_number, client_name=client_name,
            col=col, original=text, kind="manual", fix_type="broken_email",
            note="email appears duplicated (word@word@...)",
        ))
        return issues

    name_part, emails, phones = strip_email_phone(text)

    split = split_multi_person(name_part, emails, phones)
    if split:
        issues.append(CellIssue(
            row=row, client_number=client_number, client_name=client_name,
            col=col, original=text, kind="split",
            fix_type=f"{len(split)}_people",
            split_parts=split,
        ))
        return issues

    cleaned_name, applied = apply_tail_strippers(name_part)
    if applied:
        rebuilt = cleaned_name
        if emails:
            rebuilt += f" {emails[0]}"
        if phones:
            rebuilt += f" {phones[0]}"
        issues.append(CellIssue(
            row=row, client_number=client_number, client_name=client_name,
            col=col, original=text, kind="single_fix",
            fix_type=" + ".join(applied),
            cleaned=rebuilt.strip(),
        ))
        return issues

    # Suspicious long name: 4+ tokens, no separator, no credentials found
    tokens = cleaned_name.split()
    if len(tokens) >= 4:
        issues.append(CellIssue(
            row=row, client_number=client_number, client_name=client_name,
            col=col, original=text, kind="manual",
            fix_type="long_name_review",
            note=f"{len(tokens)} name tokens — verify one person vs two",
        ))

    return issues


# ── Main scan ───────────────────────────────────────────────────────────
def scan(excel_path: str,
         only_client: Optional[str]) -> tuple[list[CellIssue], list[CellIssue]]:
    wb = load_workbook(excel_path, data_only=True, read_only=True, keep_vba=False)
    if CLIENT_LIST_SHEET not in wb.sheetnames:
        raise RuntimeError(
            f"No '{CLIENT_LIST_SHEET}' sheet in workbook. "
            f"Available: {wb.sheetnames}"
        )
    ws = wb[CLIENT_LIST_SHEET]

    def col_i(c: str) -> int:
        c = c.upper()
        if len(c) == 1:
            return ord(c) - 65
        return (ord(c[0]) - 64) * 26 + (ord(c[1]) - 65)

    num_i  = col_i(COL_NUMBER)
    name_i = col_i(COL_NAME)
    contact_idx = [(c, col_i(c)) for c in CONTACT_COLS]

    cell_issues: list[CellIssue] = []
    dup_issues:  list[CellIssue] = []

    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r < DATA_START_ROW:
            continue
        number = _stringify_number(row[num_i] if num_i < len(row) else None)
        cname  = str(row[name_i] or "").strip() if name_i < len(row) else ""
        if not (number or cname):
            continue
        if only_client and number != only_client:
            continue

        seen_names: dict[str, list[str]] = {}
        for col_letter, ci in contact_idx:
            if ci >= len(row):
                continue
            cell = row[ci]
            cell_issues.extend(
                analyze_cell(cell, r, col_letter, number, cname)
            )
            if cell is None:
                continue
            text = str(cell).strip()
            if not text:
                continue
            name_part, _, _ = strip_email_phone(text)
            norm = name_part.lower().strip()
            if norm:
                seen_names.setdefault(norm, []).append(col_letter)

        for nm, cols in seen_names.items():
            if len(cols) >= 2:
                dup_issues.append(CellIssue(
                    row=r, client_number=number, client_name=cname,
                    col=", ".join(cols), original=nm, kind="duplicate",
                    fix_type="duplicate_in_client",
                    note=f"appears in {len(cols)} cells: {', '.join(cols)}",
                ))

    return cell_issues, dup_issues


# ── Report writer ───────────────────────────────────────────────────────
def write_report(cell_issues: list[CellIssue],
                 dup_issues: list[CellIssue],
                 out_path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)   # drop default sheet

    HDR_FONT  = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    HDR_FILL  = PatternFill("solid", fgColor="305496")
    BODY_FONT = Font(name="Arial", size=10)
    WRAP      = Alignment(wrap_text=True, vertical="top")

    def init_sheet(name: str, headers: list[str], widths: list[int]):
        ws = wb.create_sheet(name)
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=i, value=h)
            c.font = HDR_FONT
            c.fill = HDR_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 24
        return ws

    def format_body(ws):
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = BODY_FONT
                c.alignment = WRAP

    singles = [i for i in cell_issues if i.kind == "single_fix"]
    splits  = [i for i in cell_issues if i.kind == "split"]
    manuals = [i for i in cell_issues if i.kind == "manual"]

    # Summary first
    ws_sum = init_sheet("Summary",
        ["Category", "Count", "Description"],
        [32, 10, 78])
    ws_sum.append(["Single-cell fixes", len(singles),
                   "Strip trailing credentials, labels, apostrophes. "
                   "Low-risk. Cleaned value shown on 'Single-cell fixes' sheet."])
    ws_sum.append(["Multi-person splits", len(splits),
                   "Cells containing 2+ people. Each split becomes two cells "
                   "in adjacent Contact columns. Review each on 'Multi-person splits' sheet."])
    ws_sum.append(["Manual review", len(manuals),
                   "Broken emails, or 4+ token names with no separator that "
                   "might be two people. Needs human eyes — see 'Manual review' sheet."])
    ws_sum.append(["Duplicates within client", len(dup_issues),
                   "Same contact name in multiple Contact cells for one Client. "
                   "Decide whether to delete duplicates."])
    ws_sum.append(["", "", ""])
    ws_sum.append(["TOTAL cell issues flagged",
                   len(singles) + len(splits) + len(manuals), ""])
    format_body(ws_sum)
    # Yellow-highlight the total row
    for c in ws_sum[ws_sum.max_row]:
        c.fill = PatternFill("solid", fgColor="FFF2CC")
        c.font = Font(name="Arial", size=10, bold=True)

    # Single-cell fixes
    ws1 = init_sheet("Single-cell fixes",
        ["Row", "Client #", "Client Name", "Col", "Original", "Proposed", "Fix applied"],
        [6, 10, 28, 6, 52, 52, 32])
    for it in singles:
        ws1.append([it.row, it.client_number, it.client_name, it.col,
                    it.original, it.cleaned, it.fix_type])
    format_body(ws1)

    # Multi-person splits
    ws2 = init_sheet("Multi-person splits",
        ["Row", "Client #", "Client Name", "Orig Col", "Original",
         "Person 1 Name", "P1 Email", "P1 Phone",
         "Person 2 Name", "P2 Email", "P2 Phone",
         "Additional"],
        [6, 10, 28, 8, 52, 22, 30, 15, 22, 30, 15, 32])
    for it in splits:
        parts = it.split_parts
        p1 = parts[0] if len(parts) >= 1 else {"name": "", "email": "", "phone": ""}
        p2 = parts[1] if len(parts) >= 2 else {"name": "", "email": "", "phone": ""}
        extra = "; ".join(
            f"{p.get('name','')} / {p.get('email','')} / {p.get('phone','')}"
            for p in parts[2:]
        ) if len(parts) > 2 else ""
        ws2.append([
            it.row, it.client_number, it.client_name, it.col, it.original,
            p1.get("name", ""), p1.get("email", ""), p1.get("phone", ""),
            p2.get("name", ""), p2.get("email", ""), p2.get("phone", ""),
            extra,
        ])
    format_body(ws2)

    # Manual review
    ws3 = init_sheet("Manual review",
        ["Row", "Client #", "Client Name", "Col", "Original", "Issue", "Note"],
        [6, 10, 28, 6, 52, 22, 44])
    for it in manuals:
        ws3.append([it.row, it.client_number, it.client_name, it.col,
                    it.original, it.fix_type, it.note])
    format_body(ws3)

    # Duplicates
    ws4 = init_sheet("Duplicates within client",
        ["Row", "Client #", "Client Name", "Duplicated Name", "Cells", "Note"],
        [6, 10, 28, 32, 22, 44])
    for it in dup_issues:
        ws4.append([it.row, it.client_number, it.client_name,
                    it.original, it.col, it.note])
    format_body(ws4)

    wb.save(out_path)


def _find_excel(override: Optional[str]) -> Optional[str]:
    if override:
        return override if os.path.isfile(override) else None
    for p in DEFAULT_EXCEL_PATHS:
        if os.path.isfile(p):
            return p
    return None


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description="Analyze Client List cells for cleanup")
    ap.add_argument("--excel", help="Path to Master-Business Plan.xlsm")
    ap.add_argument("--client", help="Only scan this Client Number")
    args = ap.parse_args(argv)

    excel = _find_excel(args.excel)
    if not excel:
        print("ERROR: Master-Business Plan.xlsm not found. Tried:", file=sys.stderr)
        for p in DEFAULT_EXCEL_PATHS:
            print(f"  {p}", file=sys.stderr)
        print("Pass --excel to specify an explicit path.", file=sys.stderr)
        return 1

    print(f"Scanning: {excel}")
    try:
        cell_issues, dup_issues = scan(excel, args.client)
    except PermissionError:
        print("ERROR: Excel is open in another app (or OneDrive-locked). "
              "Close the workbook in Excel and re-run.", file=sys.stderr)
        return 1

    singles = sum(1 for i in cell_issues if i.kind == "single_fix")
    splits  = sum(1 for i in cell_issues if i.kind == "split")
    manuals = sum(1 for i in cell_issues if i.kind == "manual")

    print()
    print("Findings")
    print("─" * 40)
    print(f"  Single-cell fixes:     {singles}")
    print(f"  Multi-person splits:   {splits}")
    print(f"  Manual review needed:  {manuals}")
    print(f"  Duplicates in client:  {len(dup_issues)}")
    print("─" * 40)

    write_report(cell_issues, dup_issues, OUTPUT_PATH)
    print(f"\nReport written: {OUTPUT_PATH}")
    print("Open the report, review each sheet, then tell me which categories")
    print("to auto-apply in the 'apply' step.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
