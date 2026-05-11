"""
Excel -> comp cache enrichment.

Reads Jobs 25KS and Jobs 26KS sheets from Master-Business Plan.xlsm, normalizes
fees, collapses .1/.2/.3 add-service rows into the base project, and merges the
data into comps_cache.json. Never overwrites existing populated fields ("fill
blanks only" policy).

Called by the excel-comps-sync scheduled task. Safe to re-run; idempotent.

Returns a dict: {
    "excel_mtime": ...,
    "base_jobs_processed": N,
    "enriched": N,        # cache records whose empty fields got filled
    "added": N,           # brand-new comps added to cache from Excel
    "skipped_mtime_unchanged": bool,
}
"""
import os
import sys
import json
import warnings
from collections import defaultdict
from openpyxl import load_workbook

warnings.filterwarnings("ignore")

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _SCRIPT_DIR)
from fee_memo import write_comps_cache  # noqa

# --- Paths ---
ONEDRIVE_ROOT = os.path.dirname(_SCRIPT_DIR)
# The Business Plan folder is a SIBLING of _Projects on OneDrive.
# Expected mount path: <session>/mnt/Business Plan/Master-Business Plan.xlsm
EXCEL_PATH_CANDIDATES = [
    os.path.join(os.path.dirname(ONEDRIVE_ROOT), "Business Plan", "Master-Business Plan.xlsm"),
    "/sessions/admiring-friendly-brown/mnt/Business Plan/Master-Business Plan.xlsm",
]

CACHE_PATH = os.path.join(_SCRIPT_DIR, "comps_cache.json")
MTIME_STATE_PATH = os.path.join(_SCRIPT_DIR, "excel_sync_state.json")


# ---------------------------------------------------------------------------
# Sheet-specific column mappings (row 5 headers)
# ---------------------------------------------------------------------------
HEADERS_26KS = {
    "B": "Job", "D": "Y", "E": "Client", "F": "Name",
    "J": "SD", "K": "DD", "L": "CD",
    "M": "AS1", "N": "AS2", "O": "AS3",
    "P": "Bid_Permit", "Q": "Reimb",
    "R": "Design", "S": "CA",
    "V": "TOTAL_SIGNED", "X": "UNSIGNED",
}
HEADERS_25KS = {
    "B": "Job", "C": "Y", "D": "Client", "E": "Name",
    "J": "SD", "K": "DD", "L": "CD",
    "M": "Bid_Permit", "N": "Reimb",
    "O": "Design", "P": "CA",
    "R": "TOTAL_SIGNED", "T": "UNSIGNED",
}


def _col_to_index(letter):
    letter = letter.upper()
    if len(letter) == 1:
        return ord(letter) - 64
    return (ord(letter[0]) - 64) * 26 + (ord(letter[1]) - 64)


def _to_num(v):
    try:
        return float(v) if v not in (None, "") else 0
    except (TypeError, ValueError):
        return 0


# Markers that flag a row as a test / sandbox / placeholder entry. Any row
# whose Name contains one of these (case-insensitive) gets dropped before
# aggregation. Used to prevent test pages from polluting fee comp stats.
_TEST_MARKERS = (
    "_test",
    "automation test",
    "n8n sandbox",
    "sandbox",
    "placeholder",
    "do not use",
    "delete me",
)


def _is_meaningful_row(rec):
    """Return True if a row has enough info to be a real comp entry.

    Skips rows where:
      - Name is blank
      - Name has no alphabetic characters (e.g., "25217" by itself, or
        "26106000" — just project numbers with no scope description)
      - Name contains a test/sandbox marker (see _TEST_MARKERS)

    Real comp rows always have a project name with some descriptive text
    even if fees haven't been entered yet (a "lost" entry still has a
    client/project name attached).
    """
    name = (rec.get("Name") or "").strip()
    if not name:
        return False
    if not any(ch.isalpha() for ch in name):
        return False
    low = name.lower()
    if any(marker in low for marker in _TEST_MARKERS):
        return False
    return True


def extract_sheet(ws, col_map):
    recs = []
    for row_idx in range(6, ws.max_row + 1):
        rec = {}
        for col_letter, key in col_map.items():
            rec[key] = ws.cell(row_idx, _col_to_index(col_letter)).value
        if rec.get("Job"):
            recs.append(rec)
    return recs


def _infer_type(name):
    low = (name or "").lower()
    if any(k in low for k in ("pemb", "metal bldg", "rapid set")):
        return "PEMB"
    if any(k in low for k in ("apartment", "apartments", "multifamily")):
        return "Multifamily"
    if any(k in low for k in ("hotel", "inn", "suites", "marriot", "hilton", "cambria")):
        return "Commercial Remodel"
    if "residence" in low or "casita" in low or "adu" in low:
        return "Residential Remodel" if "remodel" in low or "addition" in low else "New Residence"
    if any(k in low for k in ("restaurant", "dunkin", "starbucks", "pizza", "cafe",
                               "bar", "brewery", "office", "dental", "retail",
                               "store", "gym", "fitness")):
        return "Tenant Improvement"
    if any(k in low for k in ("remodel", "renovation")):
        return "Commercial Remodel"
    if any(k in low for k in ("observation", "review", "repair")):
        return "Structural Observation"
    if any(k in low for k in ("new", "ground up", "commercial", "industrial",
                               "building", "bldg")):
        return "Commercial New"
    return "Other"


def find_excel():
    for p in EXCEL_PATH_CANDIDATES:
        if os.path.isfile(p):
            return p
    return None


def load_state():
    if os.path.isfile(MTIME_STATE_PATH):
        try:
            with open(MTIME_STATE_PATH) as f:
                return json.load(f)
        except Exception:
            pass
    return {}


def save_state(state):
    with open(MTIME_STATE_PATH, "w") as f:
        json.dump(state, f, indent=2)


def run(force=False):
    """
    Perform the enrichment. Returns a stats dict.
    If the Excel's mtime matches the last-seen mtime, skips work unless force=True.
    """
    excel_path = find_excel()
    if not excel_path:
        return {"error": "Master-Business Plan.xlsm not found. Is the Business Plan folder mounted?"}

    mtime = os.path.getmtime(excel_path)
    state = load_state()
    if not force and state.get("last_mtime") == mtime:
        return {
            "skipped_mtime_unchanged": True,
            "excel_mtime": mtime,
            "excel_path": excel_path,
            "enriched": 0, "added": 0,
        }

    wb = load_workbook(excel_path, data_only=True, keep_vba=False)
    excel_rows = []
    if "Jobs 25KS" in wb.sheetnames:
        excel_rows += extract_sheet(wb["Jobs 25KS"], HEADERS_25KS)
    if "Jobs 26KS" in wb.sheetnames:
        excel_rows += extract_sheet(wb["Jobs 26KS"], HEADERS_26KS)

    # Aggregate by base job (collapse .1/.2/.3). Filter out placeholder /
    # test / hollow rows up front so they never enter the comp database.
    by_base = defaultdict(lambda: {"design": 0, "cd": 0, "ca": 0, "total_signed": 0,
                                    "won": False, "name": None, "rows": 0})
    skipped_rows = 0
    for r in excel_rows:
        job = str(r.get("Job") or "").strip()
        if not job:
            continue
        if not _is_meaningful_row(r):
            skipped_rows += 1
            continue
        base = job.split(".")[0]
        agg = by_base[base]
        agg["rows"] += 1
        if not agg["name"]:
            agg["name"] = (r.get("Name") or "").strip()
        agg["design"] += _to_num(r.get("Design"))
        agg["cd"] += _to_num(r.get("CD"))
        agg["ca"] += _to_num(r.get("CA"))
        agg["total_signed"] += _to_num(r.get("TOTAL_SIGNED"))
        if (str(r.get("Y") or "")).upper().startswith("Y"):
            agg["won"] = True

    # Load current cache
    if os.path.isfile(CACHE_PATH):
        with open(CACHE_PATH) as f:
            cache = json.load(f)
    else:
        cache = {"comps": []}
    cache_by_base = {}
    for c in cache["comps"]:
        bj = str(c.get("job", "")).split(".")[0]
        if bj:
            cache_by_base[bj] = c

    enriched = 0
    added = 0
    for base, agg in by_base.items():
        total = agg["total_signed"] if agg["total_signed"] > 0 else None
        if base in cache_by_base:
            c = cache_by_base[base]
            did_enrich = False
            for field, ex_val in (("design", agg["design"]), ("cd", agg["cd"]), ("ca", agg["ca"])):
                if (c.get(field) in (None, 0)) and ex_val > 0:
                    c[field] = ex_val
                    did_enrich = True
            if total and not (c.get("total_signed") or c.get("total_proposed")):
                if agg["won"]:
                    c["total_signed"] = total
                else:
                    c["total_proposed"] = total
                did_enrich = True
            if agg["won"] and not c.get("won"):
                c["won"] = True
                did_enrich = True
            if did_enrich:
                enriched += 1
        else:
            comp = {
                "name": (agg["name"] or base)[:80],
                "job": base,
      