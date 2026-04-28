"""
Fee Analysis Memo generator.
Takes a project intake (dict) and produces a .docx memo with:
  - Project summary
  - Closest comps from the comp database (with win/loss status)
  - Fee distribution stats
  - Close-rate context and pricing tier recommendations
  - External reference points

Uses real Jobs 26KS data pulled from Notion.
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import date, datetime, timezone
import os
import json
import statistics as stats

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_COMPS_CACHE_PATH = os.path.join(_SCRIPT_DIR, "comps_cache.json")


def get_comps():
    """
    Return the active comp list. Preference order:
      1. comps_cache.json — refreshed from live Notion by scheduled task
      2. hardcoded COMPS fallback below

    The cache is a JSON object: {"generated_at": ISO8601, "comps": [...]}
    where each comp dict matches the same shape as COMPS entries.
    """
    if os.path.exists(_COMPS_CACHE_PATH):
        try:
            with open(_COMPS_CACHE_PATH, "r", encoding="utf-8") as f:
                payload = json.load(f)
            comps = payload.get("comps")
            if isinstance(comps, list) and comps:
                return comps
        except Exception:
            pass
    return COMPS


def write_comps_cache(comps, source="manual"):
    """
    Write a fresh comp list to the cache. Called by the Claude scheduled task
    after querying Jobs 26KS + Projects. `source` is just a label for logs.
    """
    payload = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": source,
        "count": len(comps),
        "comps": comps,
    }
    with open(_COMPS_CACHE_PATH, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2)
    return _COMPS_CACHE_PATH


# =====================================================================
# Comp Database — real Jobs 26KS records (pulled from Notion 2026-04-20)
# Project Type inferred from project name where Notion field not set.
# =====================================================================
COMPS = [
    # approx_sf values are rough estimates for MVP; these should be replaced
    # with actual values from the new "Approx. Structural SF" Notion field
    # as teams populate it. None = unknown (shown as "—" in memo).
    # --- PEMB / Commercial / Industrial ---
    {"name": "Faith PEMB Final Revisions",       "job": "25052001.1", "type": "PEMB",
     "design": 23600, "cd": 1100, "ca": 2300, "total_signed": 23600, "won": True,
     "partner": "JB", "approx_sf": 15000},
    {"name": "Emerald Industrial Building Suite 5", "job": "25101049", "type": "Commercial Remodel",
     "design": 1200, "cd": 1200, "ca": 300, "total_signed": 1200, "won": True,
     "partner": "JB", "approx_sf": 2500},
    {"name": "Emerald Industrial Building Suite 2", "job": "26042049", "type": "Commercial Remodel",
     "design": 1500, "cd": 1500, "ca": 400, "total_signed": 1500, "won": True,
     "partner": "CV", "approx_sf": 3000},
    {"name": "91st and McDowell Commercial Pad",  "job": "26054049", "type": "Commercial New",
     "design": 19800, "cd": 19800, "ca": 3700, "total_proposed": 23500, "won": False,
     "partner": None, "approx_sf": 10000},
    {"name": "Sunstate Companies PEMB Foundation","job": "25192033", "type": "PEMB",
     "design": 7400, "cd": 0, "ca": 850, "total_proposed": 8250, "won": False,
     "partner": None, "approx_sf": 8000},
    {"name": "Flynn Family Farm PEMB",            "job": "26073136", "type": "PEMB",
     "design": None, "cd": None, "ca": None, "total_proposed": None, "won": False,
     "partner": None, "approx_sf": None},
    {"name": "Azteca Milling Pump Building",      "job": "26007057", "type": "Commercial New",
     "design": 3800, "cd": 3800, "ca": 0, "total_proposed": 3800, "won": False,
     "partner": None, "approx_sf": 1500},
    {"name": "Apache Building Apartments",        "job": "25207019", "type": "Multifamily",
     "design": None, "cd": None, "ca": None, "total_proposed": None, "won": False,
     "partner": None, "approx_sf": None},

    # --- Residential ---
    {"name": "Coder Residence",                   "job": "25055054", "type": "New Residence",
     "design": 11600, "cd": 11600, "ca": 0, "total_signed": 11600, "won": True,
     "partner": "CV", "approx_sf": 3500},
    {"name": "Rios Remodel",                      "job": "25097072", "type": "Residential Remodel",
     "design": 3222, "cd": 3200, "ca": 140, "total_signed": 3222, "won": True,
     "partner": "CV", "approx_sf": 2000},
    {"name": "Ayala Residence Remodel",           "job": "26009051", "type": "Residential Remodel",
     "design": 1000, "cd": 1000, "ca": 0, "total_signed": 1000, "won": True,
     "partner": "CV", "approx_sf": 1500},
    {"name": "Salas Residence New ADU",           "job": "26008094", "type": "New Residence",
     "design": 4500, "cd": 0, "ca": 0, "total_proposed": 4500, "won": False,
     "partner": None, "approx_sf": 800},

    # --- Observation / Feasibility ---
    {"name": "Legacy Hotel Structural Observation","job": "25202004", "type": "Structural Observation",
     "design": 3000, "cd": 1500, "ca": 0, "total_signed": 3000, "won": True,
     "partner": "CV", "approx_sf": None},  # Observation — SF not meaningful

    # ====================================================================
    # Expansion batch — commercial / hotel / TI / retail comps pulled from
    # Jobs 26KS on 2026-04-21.  Rough SF estimates — refine as teams update.
    # ====================================================================

    # --- Tenant Improvement / Retail / Restaurant TI ---
    {"name": "PHNX Retail at Miller and Southern", "job": "240040", "type": "Tenant Improvement",
     "design": 6950, "cd": 6500, "ca": 400, "total_signed": 6950, "won": True,
     "partner": "JB", "approx_sf": 3500},
    {"name": "Roma Pizzeria Remodel", "job": "25133033", "type": "Tenant Improvement",
     "design": 1994, "cd": 1950, "ca": 0, "total_signed": 1994, "won": True,
     "partner": "CV", "approx_sf": 2000},
    {"name": "Dental Office Roanoke TX", "job": "26046049", "type": "Tenant Improvement",
     "design": 4250, "cd": 4250, "ca": 0, "total_proposed": 4250, "won": False,
     "partner": None, "approx_sf": 2500},
    {"name": "Overton Office", "job": "26026061", "type": "Tenant Improvement",
     "design": 11900, "cd": 11000, "ca": 1100, "total_proposed": 11900, "won": False,
     "partner": None, "approx_sf": 5000},

    # --- Commercial Remodel / Hotel ---
    {"name": "Wailea Hotel Remodel", "job": "240044", "type": "Commercial Remodel",
     "design": 8480, "cd": 4240, "ca": 2120, "total_signed": 8480, "won": True,
     "partner": "JB", "approx_sf": 5000},
    {"name": "Cambria Hotel VE Revisions", "job": "240038.2", "type": "Commercial Remodel",
     "design": 2400, "cd": 2400, "ca": 1225, "total_signed": 2400, "won": True,
     "partner": "CV", "approx_sf": 2000},
    {"name": "Williams Four Points Hotel Conversion", "job": "25007029", "type": "Commercial Remodel",
     "design": 14650, "cd": 10000, "ca": 150, "total_signed": 14650, "won": True,
     "partner": "JB", "approx_sf": 30000},
    {"name": "Townplace Suites Hotel Phoenix AZ", "job": "25193107", "type": "Commercial Remodel",
     "design": 47600, "cd": 45800, "ca": 4600, "total_proposed": 47600, "won": False,
     "partner": None, "approx_sf": 60000},
    {"name": "United Cambria Hotel Mesa", "job": "26018004", "type": "Commercial Remodel",
     "design": 1000, "cd": 1000, "ca": 0, "total_signed": 1000, "won": True,
     "partner": "CV", "approx_sf": 500},
]


# Which types are "close" to each other for fallback matching
TYPE_GROUPS = {
    "PEMB":                   ["PEMB", "Commercial New"],
    "Commercial New":         ["Commercial New", "PEMB"],
    "Commercial Remodel":     ["Commercial Remodel", "Commercial New"],
    "Multifamily":            ["Multifamily", "Commercial New"],
    "New Residence":          ["New Residence", "Residential Remodel"],
    "Residential Remodel":    ["Residential Remodel", "New Residence"],
    "Structural Observation": ["Structural Observation", "Peer Review", "Feasibility"],
    "Feasibility":            ["Feasibility", "Structural Observation"],
    "Peer Review":            ["Peer Review", "Structural Observation"],
    "Retrofit":               ["Retrofit", "Commercial Remodel"],
    "Tenant Improvement":     ["Tenant Improvement", "Commercial Remodel", "Commercial New"],
}


def find_comps(project_type, comps=None, min_with_fee=4, subject_sf=None):
    if comps is None:
        comps = get_comps()
    """
    Return (primary, related) comps.

    Primary = exact type match.
    Related = adjacent types added when any of these is true:
      1. Primary doesn't have at least min_with_fee comps with fee data
      2. Subject project is much larger than any primary comp (subject_sf >
         2x the max primary SF) — small primary comps aren't useful for a
         much bigger project in the same category (e.g. 24K SF driving range
         classified as TI vs. 2-5K SF pizzeria TIs).

    This keeps the comp set rich enough for reasonable statistics AND keeps
    the engine honest about when it needs broader reference points.
    """
    primary_types = TYPE_GROUPS.get(project_type, [project_type])
    primary = [c for c in comps if c["type"] == project_type]
    primary_with_fee = sum(1 for c in primary if comp_fee(c) is not None)

    # Scale mismatch check: is the subject project dramatically bigger than
    # any primary comp in its category?
    scale_mismatch = False
    if subject_sf:
        primary_sfs = [c.get("approx_sf") for c in primary if c.get("approx_sf")]
        if primary_sfs and subject_sf > max(primary_sfs) * 2:
            scale_mismatch = True

    if primary_with_fee >= min_with_fee and not scale_mismatch:
        return primary, []
    related = [c for c in comps if c["type"] in primary_types and c not in primary]
    return primary, related


def comp_fee(c):
    """Return the relevant fee for a comp (signed if won, proposed if lost)."""
    if c["won"]:
        return c.get("total_signed")
    return c.get("total_proposed")


def analyze_comps(comps):
    """Compute statistics on won fees and win rate."""
    won = [c for c in comps if c["won"] and comp_fee(c) is not None]
    lost = [c for c in comps if not c["won"] and comp_fee(c) is not None]
    won_fees = [comp_fee(c) for c in won]

    stats_d = {}
    if won_fees:
        won_fees_sorted = sorted(won_fees)
        stats_d["won_count"] = len(won_fees)
        stats_d["lost_count"] = len(lost)
        stats_d["median"] = stats.median(won_fees_sorted)
        stats_d["min"] = min(won_fees_sorted)
        stats_d["max"] = max(won_fees_sorted)
        n = len(won_fees_sorted)
        if n >= 4:
            stats_d["p25"] = won_fees_sorted[int(n * 0.25)]
            stats_d["p75"] = won_fees_sorted[int(n * 0.75)]
        else:
            stats_d["p25"] = stats_d["min"]
            stats_d["p75"] = stats_d["max"]
    else:
        stats_d = {"won_count": 0, "lost_count": len(lost), "median": None,
                   "min": None, "max": None, "p25": None, "p75": None}

    # Win rate in category
    total = len(won) + len(lost)
    stats_d["win_rate"] = (len(won) / total * 100) if total else None
    stats_d["total_comps"] = total
    return stats_d


def recommend_tiers(stats_d, project_type):
    """Compute Conservative / Market / Stretch fee tiers from stats."""
    if not stats_d["median"]:
        return None

    median = stats_d["median"]
    p75 = stats_d["p75"]
    max_won = stats_d["max"]

    # Conservative = median of wins (preserves current close behavior)
    # Market     = max(median * 1.25, P75 of wins) — tests upward price discipline
    # Stretch    = max(market * 1.20, max_won * 1.30) — always strictly above market
    conservative = int(round(median / 100) * 100)
    market_raw = max(median * 1.25, p75)
    market = int(round(market_raw / 100) * 100)
    stretch_raw = max(market_raw * 1.20, max_won * 1.30)
    stretch = int(round(stretch_raw / 100) * 100)

    return {
        "conservative": conservative,
        "market": market,
        "stretch": stretch,
    }


# =====================================================================
# Document rendering helpers
# =====================================================================
def _set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def _set_cell_borders(cell, color='888888'):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), color)
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def _fill_row(table, row_idx, values, bold=False, bg=None, font_size=10, white_text=False):
    for i, v in enumerate(values):
        cell = table.rows[row_idx].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        run = p.add_run(str(v))
        run.font.size = Pt(font_size)
        run.font.bold = bold
        if white_text:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if bg:
            _set_cell_bg(cell, bg)
        _set_cell_borders(cell)


def _add_heading(doc, text, size=13, color=None, space_before=14, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = True
    if color:
        run.font.color.rgb = color


def _add_para(doc, text, size=11, italic=False, bold=False, space_after=6, space_before=0, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if space_before:
        p.paragraph_format.space_before = Pt(space_before)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.italic = italic
    run.font.bold = bold


def _add_bullet(doc, text, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run('•  ' + text)
    r.font.size = Pt(size)


def _fmt(v):
    if v is None:
        return "—"
    return f"${v:,}"


def _fmt_psf(fee, sf):
    """Format $/SF, return "—" if either is missing."""
    if not fee or not sf:
        return "—"
    return f"${fee/sf:.2f}/SF"


def _fmt_sf(sf):
    if not sf:
        return "—"
    return f"{int(sf):,} SF"


# =====================================================================
# Main memo render
# =====================================================================
def render_fee_memo(project, out_dir, overall_close_rate=70, filename=None):
    """
    project dict keys:
        project_name (short, e.g., 'General Metal Construction PEMB')
        ks_job_number
        project_type (must match keys in TYPE_GROUPS)
        approx_sf (number, optional)
        jurisdiction
        location (e.g., 'Phoenix, AZ')
        scope_description
        partner (JB/CV/GO)
        client_company_name
        engineering_status
    """
    doc = Document()

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)

    KS_BLUE = RGBColor(0x1F, 0x3A, 0x5F)

    # -------- Title --------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('FEE ANALYSIS MEMO')
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = KS_BLUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Internal Working Document  —  KINGDOM STRUCTURAL LLC')
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    p = doc.add_paragraph()
    r = p.add_run('_' * 80)
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # -------- Metadata table --------
    meta = doc.add_table(rows=5, cols=2)
    meta.columns[0].width = Inches(1.8)
    meta.columns[1].width = Inches(4.7)
    partner_map = {"JB": "Jonathan Brusnahan, SE, PE", "CV": "Chris Valdez, PE", "GO": "Griffin O'Reilly"}
    meta_rows = [
        ('Project', project.get('project_name', '')),
        ('KS Job Number', project.get('ks_job_number', '')),
        ('Project Type', project.get('project_type', '')),
        ('Prepared', date.today().strftime("%B %d, %Y")),
        ('Lead Partner', partner_map.get(project.get('partner', ''), project.get('partner', ''))),
    ]
    for i, (k, v) in enumerate(meta_rows):
        meta.rows[i].cells[0].text = ''
        pp = meta.rows[i].cells[0].paragraphs[0]
        rr = pp.add_run(k)
        rr.font.bold = True
        rr.font.size = Pt(10)
        meta.rows[i].cells[1].text = ''
        pp = meta.rows[i].cells[1].paragraphs[0]
        rr = pp.add_run(str(v))
        rr.font.size = Pt(10)

    # -------- 1. Project Summary --------
    _add_heading(doc, '1. Project Summary', color=KS_BLUE, space_before=18)
    subject_sf = project.get('approx_sf')
    sf_note = f"Approx. {subject_sf:,} SF structural. " if subject_sf else ""
    loc_note = f"Location: {project.get('location', 'TBD')}. "
    jurisdiction_note = f"Jurisdiction: {project.get('jurisdiction', 'TBD')}. "
    _add_para(doc,
        project.get('scope_description', 'Scope not yet set in Notion.') + ' ' +
        sf_note + loc_note + jurisdiction_note,
        space_after=4
    )

    # Gather comps for all subsequent sections
    ptype = project.get('project_type', '')
    primary, related = find_comps(ptype, subject_sf=subject_sf)
    all_comps = primary + related
    lost_with_fee = [c for c in all_comps if not c["won"] and comp_fee(c) is not None]
    won_with_fee = [c for c in all_comps if c["won"] and comp_fee(c) is not None]

    # -------- 2. LOSS SIGNALS (prominent early section) --------
    _add_heading(doc, '2. Loss Signals — Where Clients Walked', color=RGBColor(0xB0, 0x3A, 0x2E),
                 space_before=14)
    if lost_with_fee:
        # Cap loss-signals display at 10 most-relevant (highest proposed fee first —
        # these are the most informative ceiling references)
        MAX_LOSS_DISPLAY = 10
        displayed_losses = sorted(lost_with_fee, key=lambda c: -(comp_fee(c) or 0))[:MAX_LOSS_DISPLAY]

        lead = f'{len(lost_with_fee)} proposal(s) in or near this project class did not close. '
        if len(lost_with_fee) > MAX_LOSS_DISPLAY:
            lead += f'Showing the {MAX_LOSS_DISPLAY} highest-proposed (all {len(lost_with_fee)} inform the stats below). '
        lead += 'These are ceiling references — prices at which clients chose not to sign. Use them to calibrate the upper edge of your pricing tiers.'
        _add_para(doc, lead, size=10, italic=True, space_after=6)

        loss_table = doc.add_table(rows=len(displayed_losses) + 1, cols=5)
        widths = [Inches(2.3), Inches(1.2), Inches(1.0), Inches(1.0), Inches(1.0)]
        for i, w in enumerate(widths):
            loss_table.columns[i].width = w
        _fill_row(loss_table, 0, ['Project', 'Type', 'Proposed Fee', 'SF', '$/SF'],
                  bold=True, bg='B03A2E', white_text=True, font_size=10)
        for i, c in enumerate(displayed_losses, start=1):
            fee = comp_fee(c)
            sf = c.get("approx_sf")
            _fill_row(loss_table, i, [
                c['name'][:40],
                c['type'],
                _fmt(fee),
                _fmt_sf(sf),
                _fmt_psf(fee, sf),
            ], bg='FFF2F2', font_size=9)

        # Highlight dangerous proximity: lost bids near or below where we'd recommend pricing
        if won_with_fee:
            lowest_lost = min(comp_fee(c) for c in lost_with_fee)
            highest_won = max(comp_fee(c) for c in won_with_fee)
            if lowest_lost <= highest_won:
                _add_para(doc,
                    f'⚠ Notable: lowest lost bid ({_fmt(lowest_lost)}) sits at or below the highest won '
                    f'({_fmt(highest_won)}) in this category. Price alone does not determine close — '
                    f'scope fit, client relationship, and timing matter. Treat these as directional signals, '
                    f'not hard ceilings.',
                    size=10, bold=True, space_before=6, space_after=4
                )
    else:
        _add_para(doc,
            'No lost proposals recorded in this category yet. Pricing ceiling is unknown — '
            'recommend treating Market tier as aspirational and watching closely for the first loss.',
            size=10, italic=True)

    # -------- 3. Comparable Projects --------
    _add_heading(doc, '3. Comparable Projects from Jobs 26KS', color=KS_BLUE)

    # Select up to 10 most-relevant comps for display (stats still use ALL comps)
    # Ranking: exact type match > related type, comps with fees > without,
    # closer SF to subject SF preferred.
    def _relevance_key(c):
        exact_type = 0 if c["type"] == ptype else 1
        has_fee = 0 if comp_fee(c) is not None else 1
        sf_distance = 0
        if subject_sf and c.get("approx_sf"):
            sf_distance = abs(c["approx_sf"] - subject_sf) / max(subject_sf, 1)
        else:
            sf_distance = 999  # unknown SF = deprioritized for display
        return (exact_type, has_fee, sf_distance)

    MAX_DISPLAY = 10
    ranked_comps = sorted(all_comps, key=_relevance_key)
    display_comps = ranked_comps[:MAX_DISPLAY]

    lead_text = f'Primary match: project type "{ptype}". '
    if not primary:
        lead_text += 'No exact-type comps yet. '
    if related:
        lead_text += f'{len(primary)} exact match(es) plus {len(related)} related-type. '
    if len(all_comps) > MAX_DISPLAY:
        lead_text += f'Showing the {MAX_DISPLAY} most relevant (ranked by type match, fee availability, and SF proximity); all {len(all_comps)} are used in the statistics below.'
    _add_para(doc, lead_text, size=10, italic=True, space_after=4)

    comp_table = doc.add_table(rows=len(display_comps) + 1, cols=6)
    widths = [Inches(2.0), Inches(1.2), Inches(0.9), Inches(0.9), Inches(0.9), Inches(0.7)]
    for i, w in enumerate(widths):
        comp_table.columns[i].width = w
    _fill_row(comp_table, 0, ['Project', 'Type', 'Total Fee', 'SF', '$/SF', 'Outcome'],
              bold=True, bg='1F3A5F', white_text=True, font_size=10)
    for i, c in enumerate(display_comps, start=1):
        outcome = 'WON' if c['won'] else 'LOST'
        fee = comp_fee(c)
        sf = c.get("approx_sf")
        bg = None
        if outcome == 'LOST':
            bg = 'FFF2F2'
        elif i % 2 == 0:
            bg = 'F5F5F5'
        _fill_row(comp_table, i, [
            c['name'][:36],
            c['type'],
            _fmt(fee),
            _fmt_sf(sf),
            _fmt_psf(fee, sf),
            outcome,
        ], bg=bg, font_size=9)

    # Note on SF data coverage (of ALL comps, not just displayed)
    sf_coverage = sum(1 for c in all_comps if c.get("approx_sf"))
    _add_para(doc,
        f'SF coverage: {sf_coverage} of {len(all_comps)} total comps have SF recorded. '
        f'SF values on older records are partner estimates and will sharpen as teams populate '
        f'the "Approx. Structural SF" field on past projects.',
        size=9, italic=True, space_before=4
    )

    # -------- 4. Fee Distribution + Win Rate Context --------
    _add_heading(doc, '4. Fee Distribution and Win Rate', color=KS_BLUE)
    stats_d = analyze_comps(all_comps)

    if stats_d["won_count"] > 0:
        dist_table = doc.add_table(rows=2, cols=5)
        for i, w in enumerate([Inches(1.3)] * 5):
            dist_table.columns[i].width = w
        _fill_row(dist_table, 0, ['Won Count', 'Lost Count', 'Median', 'P25', 'P75 / Max'],
                  bold=True, bg='1F3A5F', white_text=True, font_size=10)
        _fill_row(dist_table, 1, [
            stats_d["won_count"], stats_d["lost_count"],
            _fmt(stats_d["median"]), _fmt(stats_d["p25"]),
            f'{_fmt(stats_d["p75"])} / {_fmt(stats_d["max"])}'
        ], font_size=10)

        cat_win_rate = f'{stats_d["win_rate"]:.0f}%' if stats_d["win_rate"] else 'N/A'
        _add_para(doc,
            f'Category win rate: {cat_win_rate} ({stats_d["won_count"]} won / '
            f'{stats_d["lost_count"]} lost).    '
            f'Overall KS close rate: {overall_close_rate}%.    '
            f'Industry benchmark: 30–50%.',
            size=10, space_before=6, space_after=4
        )
    else:
        _add_para(doc,
            'No won comps in category yet — pricing tiers below are based on related-type projects '
            'and industry heuristics only. Recommend broader comp matching once more category data is available.',
            size=10, italic=True)

    # -------- 5. Recommended Fee Tiers --------
    _add_heading(doc, '5. Recommended Fee Tiers', color=KS_BLUE)
    tiers = recommend_tiers(stats_d, ptype)

    if tiers:
        tier_table = doc.add_table(rows=4, cols=5)
        for i, w in enumerate([Inches(1.1), Inches(1.1), Inches(0.9), Inches(2.3), Inches(1.0)]):
            tier_table.columns[i].width = w
        _fill_row(tier_table, 0, ['Tier', 'Total Fee', '$/SF', 'Basis', 'Exp. Close'],
                  bold=True, bg='1F3A5F', white_text=True, font_size=10)

        cons_close = f'~{min(90, overall_close_rate+5)}%'
        market_close = '~50–55%'
        stretch_close = '~30–40%'

        median = stats_d["median"] or 0
        max_won = stats_d["max"] or 0

        tier_rows = [
            ['Conservative', _fmt(tiers["conservative"]),
             _fmt_psf(tiers["conservative"], subject_sf),
             f'Matches category median (${median:,}). Safe — preserves current win behavior.',
             cons_close],
            ['Market  ★', _fmt(tiers["market"]),
             _fmt_psf(tiers["market"], subject_sf),
             f'+25% over median OR at P75 of wins. Tests upward price discipline.',
             market_close],
            ['Stretch', _fmt(tiers["stretch"]),
             _fmt_psf(tiers["stretch"], subject_sf),
             f'Above highest won (${max_won:,}). Tests ceiling when scope supports premium.',
             stretch_close],
        ]
        for i, row in enumerate(tier_rows, start=1):
            bg = 'E8F0F8' if i == 2 else ('F5F5F5' if i % 2 == 0 else None)
            _fill_row(tier_table, i, row, bg=bg, font_size=9)

        # Danger-zone check: any lost bids within ±20% of each tier's total fee?
        # That's the pricing band where losses actually occurred.
        if lost_with_fee:
            for tier_name, fee_val in [('Conservative', tiers['conservative']),
                                        ('Market', tiers['market']),
                                        ('Stretch', tiers['stretch'])]:
                nearby = [c for c in lost_with_fee
                          if 0.80 * fee_val <= comp_fee(c) <= 1.20 * fee_val]
                if nearby:
                    examples = '; '.join(
                        f'{c["name"][:30]} lost at {_fmt(comp_fee(c))}'
                        for c in nearby[:2]
                    )
                    _add_para(doc,
                        f'⚠ {tier_name} tier ({_fmt(fee_val)}) is within 20% of known lost bids: {examples}. '
                        f'This is a ceiling zone — losses have occurred at this price point.',
                        size=10, bold=True, space_before=6, space_after=0
                    )

        _add_para(doc,
            '★ Market tier recommended as default posture. Calibrate over 6–8 proposals — if close rate '
            'holds above 55% at Market tier, move to Stretch as new baseline for this category.',
            size=10, italic=True, space_before=6)
    else:
        _add_para(doc, 'Unable to compute tiers — no won comps in category.',
                  size=10, italic=True)

    # -------- 6. External Reference Points --------
    _add_heading(doc, '6. External Reference Points', color=KS_BLUE)
    refs = [
        'Structural engineering fees as % of structural construction cost: typically 1.5–3.0% '
        '(higher end for complex retrofits and PEMBs, lower end for standard new-build).',
        'Phoenix metro construction demand has remained elevated through 2025–2026. '
        'Market supports upward fee pressure.',
        'Industry close-rate benchmark: ACEC/SEAoA firm surveys indicate 30–50% bid-to-close rates '
        'are typical. KS overall 70% is an outlier.',
        'Caveat: External benchmarks are rough. Own comp data will dominate this analysis as the '
        'comp database grows (currently {n} projects).'.format(n=len(get_comps())),
    ]
    for ref in refs:
        _add_bullet(doc, ref)

    # -------- 7. Partner Decision --------
    _add_heading(doc, '7. Partner Decision', color=KS_BLUE)
    _add_para(doc, 'Circle selected tier and record final fee below. Contract fee schedule will be populated with this value.',
              size=10, italic=True, space_after=8)

    if tiers:
        decision = doc.add_table(rows=1, cols=4)
        for i, w in enumerate([Inches(1.6), Inches(1.6), Inches(1.6), Inches(1.7)]):
            decision.columns[i].width = w
        _fill_row(decision, 0, [
            f'☐ Conservative  {_fmt(tiers["conservative"])}',
            f'☐ Market  {_fmt(tiers["market"])}',
            f'☐ Stretch  {_fmt(tiers["stretch"])}',
            'Final Fee: $________'
        ], bold=True, font_size=10)

    _add_para(doc, 'Signed: ______________________________  Date: _______________',
              size=10, space_before=20)

    # -------- Footer --------
    p = doc.add_paragraph()
    r = p.add_run('_' * 80)
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        f'Generated from Jobs 26KS comp database ({len(get_comps())} projects) on '
        f'{date.today().isoformat()}. Recommendations are model-generated and subject to partner override.'
    )
    r.font.size = Pt(8)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # Save
    os.makedirs(out_dir, exist_ok=True)
    if filename is None:
        filename = f'{project.get("ks_job_number","")} Fee Analysis Memo - {project.get("project_name","")}.docx'
    filename = ''.join(c for c in filename if c not in '<>:"/\\|?*')
    out_path = os.path.join(out_dir, filename)
    doc.save(out_path)
    return out_path, stats_d, tiers


# =====================================================================
# Run against real Trevor Pan project
# =====================================================================
if __name__ == "__main__":
    trevor_pan_pemb = {
        "project_name": "General Metal Construction PEMB",
        "ks_job_number": "26102144",
        "project_type": "PEMB",
        "approx_sf": 12000,
        "jurisdiction": "Phoenix",
        "location": "Phoenix, AZ",
        "scope_description": (
            "Pre-engineered metal building (PEMB) project including structural engineering, "
            "drafting, and construction administration for the primary steel frame, secondary "
            "members, and foundation design."
        ),
        "partner": "CV",
        "client_company_name": "Trevor Pan Architects",
        "engineering_status": "Proposal Sent",
    }
    out_dir = "/sessions/admiring-friendly-brown/mnt/outputs/26102144 General Metal Construction PEMB"
    out_path, stats_d, tiers = render_fee_memo(trevor_pan_pemb, out_dir)
    print(f"Rendered: {out_path}")
    print(f"\nStats: {stats_d}")
    print(f"\nTiers: {tiers}")
